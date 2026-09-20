from __future__ import annotations

import io
from pathlib import Path

from backend.app.documents.models import DocumentCell, DocumentFormat, DocumentPage, DocumentTable, UnifiedDocument
from backend.app.documents.readers.base import DocumentReader


class XlsxReader(DocumentReader):
    def can_read(self, filename: str, content_bytes: bytes | None = None) -> bool:
        return Path(filename).suffix.lower() in (".xlsx", ".xlsm", ".xltx")

    def read(self, content_bytes: bytes, filename: str, source_reference: str = "") -> UnifiedDocument:
        import openpyxl

        try:
            stream = io.BytesIO(content_bytes)
            wb = openpyxl.load_workbook(stream, data_only=True)
            text_lines: list[str] = []
            tables: list[DocumentTable] = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_rows: list[list[object]] = []
                sheet_cells: list[list[DocumentCell]] = []
                for row in ws.iter_rows():
                    row_cells: list[object] = []
                    row_locations: list[DocumentCell] = []
                    for cell in row:
                        value = cell.value
                        if value is None or (isinstance(value, str) and not value.strip()):
                            continue
                        cleaned = value.strip().replace("\n", " ") if isinstance(value, str) else value
                        row_cells.append(cleaned)
                        row_locations.append(
                            DocumentCell(
                                value=cleaned,
                                row_index=cell.row,
                                column_index=cell.column,
                                coordinate=cell.coordinate,
                            )
                        )
                    if row_cells:
                        sheet_rows.append(row_cells)
                        sheet_cells.append(row_locations)
                        if len(row_cells) >= 2:
                            text_lines.append(f"{row_cells[0]}: {' '.join(str(cell) for cell in row_cells[1:])}")
                        elif len(row_cells) == 1:
                            text_lines.append(str(row_cells[0]))

                if sheet_rows:
                    tables.append(
                        DocumentTable(
                            rows=sheet_rows,
                            page_number=1,
                            title=sheet_name,
                            cells=sheet_cells,
                        )
                    )

            full_text = "\n".join(text_lines)
            page = DocumentPage(page_number=1, text=full_text)

            return UnifiedDocument(
                raw_text=full_text,
                pages=[page],
                tables=tables,
                format=DocumentFormat.XLSX,
                reader_used="XlsxReader",
                filename=filename,
                source_reference=source_reference,
                extraction_quality=1.0 if full_text.strip() else 0.0,
                extraction_status="EXTRACTED" if full_text.strip() else "UNREADABLE",
                metadata={"sheets_count": len(wb.sheetnames)},
            )
        except Exception as exc:
            return UnifiedDocument(
                raw_text="",
                pages=[],
                tables=[],
                format=DocumentFormat.XLSX,
                reader_used="XlsxReader",
                filename=filename,
                source_reference=source_reference,
                extraction_quality=0.0,
                extraction_status="FAILED",
                error_message=str(exc),
                metadata={"error": str(exc)},
            )
