from __future__ import annotations

import io
from pathlib import Path

from backend.app.documents.models import DocumentCell, DocumentFormat, DocumentPage, DocumentTable, UnifiedDocument
from backend.app.documents.readers.base import DocumentReader


class DocxReader(DocumentReader):
    def can_read(self, filename: str, content_bytes: bytes | None = None) -> bool:
        return Path(filename).suffix.lower() == ".docx"

    def read(self, content_bytes: bytes, filename: str, source_reference: str = "") -> UnifiedDocument:
        import docx

        try:
            stream = io.BytesIO(content_bytes)
            doc = docx.Document(stream)
            text_lines: list[str] = []
            tables: list[DocumentTable] = []

            for p in doc.paragraphs:
                cleaned = p.text.strip()
                if cleaned:
                    text_lines.append(cleaned)

            for t_idx, table in enumerate(doc.tables, start=1):
                table_rows: list[list[str]] = []
                table_cells: list[list[DocumentCell]] = []
                for row_idx, row in enumerate(table.rows, start=1):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # Deduplicate adjacent duplicate cells caused by merged cells
                    deduped: list[str] = []
                    for c in cells:
                        if not deduped or c != deduped[-1]:
                            deduped.append(c)
                    if any(deduped):
                        table_rows.append(deduped)
                        table_cells.append(
                            [
                                DocumentCell(
                                    value=value,
                                    row_index=row_idx,
                                    column_index=column_idx,
                                    coordinate=f"R{row_idx}C{column_idx}",
                                )
                                for column_idx, value in enumerate(deduped, start=1)
                            ]
                        )
                        if len(deduped) >= 2:
                            text_lines.append(f"{deduped[0]}: {' '.join(deduped[1:])}")
                        elif len(deduped) == 1:
                            text_lines.append(deduped[0])
                if table_rows:
                    tables.append(
                        DocumentTable(
                            rows=table_rows,
                            page_number=1,
                            title=f"Table_{t_idx}",
                            cells=table_cells,
                        )
                    )

            full_text = "\n".join(text_lines)
            page = DocumentPage(page_number=1, text=full_text)

            return UnifiedDocument(
                raw_text=full_text,
                pages=[page],
                tables=tables,
                format=DocumentFormat.DOCX,
                reader_used="DocxReader",
                filename=filename,
                source_reference=source_reference,
                extraction_quality=1.0 if full_text.strip() else 0.0,
                extraction_status="EXTRACTED" if full_text.strip() else "UNREADABLE",
                metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
            )
        except Exception as exc:
            return UnifiedDocument(
                raw_text="",
                pages=[],
                tables=[],
                format=DocumentFormat.DOCX,
                reader_used="DocxReader",
                filename=filename,
                source_reference=source_reference,
                extraction_quality=0.0,
                extraction_status="FAILED",
                error_message=str(exc),
                metadata={"error": str(exc)},
            )
