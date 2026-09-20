from __future__ import annotations

import io
from pathlib import Path

import pytest

from backend.app.documents.models import DocumentFormat, DocumentType, UnifiedDocument
from backend.app.documents.readers.docx_reader import DocxReader
from backend.app.documents.readers.pdf_text import PdfTextReader
from backend.app.documents.readers.plain_text import PlainTextReader
from backend.app.documents.readers.xlsx_reader import XlsxReader
from backend.app.extraction.extractor import DeterministicDocumentExtractor
from backend.app.extraction.models import CanonicalField, FieldStatus, MappingMethod


BUNDLE_ATTACHMENTS = Path("data/bundle/attachments")


def _text_document(text: str, role: DocumentType = DocumentType.SI) -> UnifiedDocument:
    document = PlainTextReader().read(text.encode(), "synthetic.txt", "memory/synthetic.txt")
    document.document_type = role
    return document


def _complete_text() -> str:
    return """SHIPPING INSTRUCTION
Shipper:
Alpha Logistics Sdn. Bhd.
12 Port Road
Consignee: Beta Trading Ltd
Notify Party: Gamma Forwarding
POL: Port Klang
Port of Discharge: Singapore
No. of Containers: 6 x 40'HC
Gross Wt (kgs) (毛重 KGS): 22,000 KG
"""


@pytest.mark.req("EXT-01")
@pytest.mark.req("EXT-04")
@pytest.mark.req("EXT-05a")
def test_one_pass_text_extraction_has_exactly_seven_fields_and_preserves_raw_and_canonical_values():
    result = DeterministicDocumentExtractor().extract(_text_document(_complete_text()))

    assert set(result.fields) == set(CanonicalField)
    assert len(result.fields) == 7
    assert result.fields[CanonicalField.SHIPPER].raw_value == "Alpha Logistics Sdn. Bhd.\n12 Port Road"
    assert result.fields[CanonicalField.SHIPPER].canonical_value == "Alpha Logistics Sdn. Bhd.\n12 Port Road"
    assert result.fields[CanonicalField.PORT_OF_LOADING].raw_label == "POL"
    assert result.fields[CanonicalField.PORT_OF_LOADING].mapping_method == MappingMethod.ALIAS_DICTIONARY

    gross = result.fields[CanonicalField.GROSS_WEIGHT_KG]
    assert gross.raw_label == "Gross Wt (kgs) (毛重 KGS)"
    assert gross.raw_value == "22,000 KG"
    assert gross.canonical_value == 22000
    assert gross.mapping_method == MappingMethod.BILINGUAL_LABEL_NORMALIZATION
    assert gross.source_location.line_number == 10

    containers = result.fields[CanonicalField.CONTAINER_COUNT]
    assert containers.raw_value == "6 x 40'HC"
    assert containers.canonical_value == 6
    assert containers.evidence["container_type"] == "40'HC"


@pytest.mark.req("EXT-02")
def test_one_extractor_invocation_returns_the_complete_seven_field_result():
    class SpyExtractor:
        def __init__(self):
            self.calls = 0
            self.delegate = DeterministicDocumentExtractor()

        def extract(self, document):
            self.calls += 1
            return self.delegate.extract(document)

    extractor = SpyExtractor()
    result = extractor.extract(_text_document(_complete_text()))

    assert extractor.calls == 1
    assert set(result.fields) == set(CanonicalField)


@pytest.mark.req("MAP-04")
@pytest.mark.req("EXT-05a")
def test_deterministic_provenance_and_contextual_consignee_are_emitted_without_llm():
    document = _text_document(
        """BILL OF LADING (DRAFT)
To the Order of ABC Bank
Notify Party: Independent Notify Ltd
POL: Port Klang
POD: Singapore
Containers: 3
Gross Weight: 22000 kg
""",
        DocumentType.DRAFT_BL,
    )
    result = DeterministicDocumentExtractor().extract(document)

    consignee = result.fields[CanonicalField.CONSIGNEE]
    assert consignee.raw_label == "To the Order of"
    assert consignee.raw_value == "ABC Bank"
    assert consignee.mapping_method == MappingMethod.CONTEXTUAL_BUSINESS_RULE
    assert result.fields[CanonicalField.NOTIFY_PARTY].canonical_value == "Independent Notify Ltd"
    assert all(field.mapping_method != MappingMethod.LLM_RESOLVED for field in result.fields.values())


@pytest.mark.req("EXT-01")
@pytest.mark.req("EXT-05a")
def test_missing_conflicting_and_independent_document_values_are_not_fabricated():
    duplicate = DeterministicDocumentExtractor().extract(
        _text_document("Shipper: Alpha Ltd\nShipper: Beta Ltd\nNET WEIGHT: 18000 KG")
    )
    assert duplicate.fields[CanonicalField.SHIPPER].status == FieldStatus.AMBIGUOUS
    assert len(duplicate.fields[CanonicalField.SHIPPER].evidence["candidates"]) == 2
    assert duplicate.fields[CanonicalField.GROSS_WEIGHT_KG].status == FieldStatus.MISSING
    assert duplicate.fields[CanonicalField.GROSS_WEIGHT_KG].raw_value is None

    si = DeterministicDocumentExtractor().extract(_text_document("Gross Weight: 22000 KG"))
    bl = DeterministicDocumentExtractor().extract(
        _text_document("NET WEIGHT: 18000 KG", DocumentType.DRAFT_BL)
    )
    assert si.fields[CanonicalField.GROSS_WEIGHT_KG].canonical_value == 22000
    assert bl.fields[CanonicalField.GROSS_WEIGHT_KG].status == FieldStatus.MISSING
    for field in (
        CanonicalField.CONSIGNEE,
        CanonicalField.NOTIFY_PARTY,
        CanonicalField.PORT_OF_LOADING,
        CanonicalField.PORT_OF_DISCHARGE,
    ):
        assert si.fields[field].status == FieldStatus.MISSING
        assert bl.fields[field].status == FieldStatus.MISSING


@pytest.mark.req("EXT-05a")
def test_txt_pdf_and_docx_materializations_feed_the_same_deterministic_extractor():
    extractor = DeterministicDocumentExtractor()

    txt = PlainTextReader().read(_complete_text().encode(), "sample.txt", "memory/sample.txt")
    txt.document_type = DocumentType.SI
    assert extractor.extract(txt).fields[CanonicalField.GROSS_WEIGHT_KG].canonical_value == 22000

    pdf_path = BUNDLE_ATTACHMENTS / "email_059_SI.pdf"
    pdf = PdfTextReader().read(pdf_path.read_bytes(), pdf_path.name, str(pdf_path))
    pdf.document_type = DocumentType.SI
    pdf_result = extractor.extract(pdf)
    assert pdf.format == DocumentFormat.PDF_TEXT
    assert pdf_result.fields[CanonicalField.CONTAINER_COUNT].canonical_value == 6
    assert pdf_result.fields[CanonicalField.GROSS_WEIGHT_KG].canonical_value == 131322
    assert pdf_result.fields[CanonicalField.SHIPPER].source_location.page_number == 1

    import docx

    source = docx.Document()
    source.add_paragraph("SHIPPING INSTRUCTION")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Shipper"
    table.cell(0, 1).text = "DOCX Shipper Ltd"
    table.cell(1, 0).text = "Gross Weight"
    table.cell(1, 1).text = "22000 kg"
    payload = io.BytesIO()
    source.save(payload)
    docx_document = DocxReader().read(payload.getvalue(), "sample.docx", "memory/sample.docx")
    docx_document.document_type = DocumentType.SI
    docx_result = extractor.extract(docx_document)
    assert docx_result.fields[CanonicalField.SHIPPER].canonical_value == "DOCX Shipper Ltd"
    assert docx_result.fields[CanonicalField.SHIPPER].mapping_method == MappingMethod.TABLE_STRUCTURE
    assert docx_result.fields[CanonicalField.SHIPPER].source_location.table_name == "Table_1"


@pytest.mark.req("EXT-04")
@pytest.mark.req("EXT-05a")
@pytest.mark.req("MAP-04")
def test_xlsx_native_numerics_and_cell_coordinates_survive_extraction():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Shipment"
    sheet.append(["SHIPPING INSTRUCTION"])
    sheet.append(["Gross Weight", 22000])
    sheet.append(["Container Count", 6])
    payload = io.BytesIO()
    workbook.save(payload)

    document = XlsxReader().read(payload.getvalue(), "sample.xlsx", "memory/sample.xlsx")
    document.document_type = DocumentType.SI
    result = DeterministicDocumentExtractor().extract(document)
    gross = result.fields[CanonicalField.GROSS_WEIGHT_KG]
    containers = result.fields[CanonicalField.CONTAINER_COUNT]

    assert gross.raw_value == 22000 and isinstance(gross.raw_value, int)
    assert gross.canonical_value == 22000 and isinstance(gross.canonical_value, int)
    assert gross.mapping_method == MappingMethod.TABLE_STRUCTURE
    assert gross.source_location.table_name == "Shipment"
    assert gross.source_location.row_index == 2
    assert gross.source_location.label_cell == "A2"
    assert gross.source_location.value_cell == "B2"
    assert containers.raw_value == 6 and containers.canonical_value == 6
